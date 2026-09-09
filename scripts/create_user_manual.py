from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "user-manual"
ASSETS = ROOT.parents[1] / "deliverables" / "procureflow-video" / "assets"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = "123C69"
GREEN = "3F7A49"
PALE = "EDF4EF"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:color"), GREEN)
    borders.append(bottom)
    pPr.append(borders)

def add_picture(doc, name, caption):
    path = ASSETS / name
    if path.exists():
        doc.add_picture(str(path), width=Inches(6.65))
        p = doc.paragraphs[-1]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.style = doc.styles["Caption"]

def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
for style_name, size, color in [("Title", 34, BLUE), ("Heading 1", 22, BLUE), ("Heading 2", 15, GREEN), ("Heading 3", 12, BLUE)]:
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True

header = section.header.paragraphs[0]
header.text = "PROCUREFLOW  |  USER MANUAL  |  SIH26032"
header.runs[0].font.color.rgb = RGBColor.from_string(BLUE)
header.runs[0].font.bold = True
header.runs[0].font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("ProcureFlow · Hackathon demonstration build · Version 1.5     ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PF")
r.bold = True; r.font.size = Pt(30); r.font.color.rgb = RGBColor.from_string(GREEN)
title = doc.add_paragraph("ProcureFlow User Manual", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Farmer procurement scheduling, live queue management and direct support")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(15)
subtitle.runs[0].font.color.rgb = RGBColor.from_string(GREEN)
meta = doc.add_paragraph("Version 1.5  •  10 September 2026  •  Software prototype for SIH26032")
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_picture(doc, "01_title.png", "ProcureFlow — a focused, multilingual procurement-service MVP")
doc.add_paragraph("This manual covers the public website, installable Android web app, automated ProcureBot assistant, farmer/staff-to-admin live chat, notifications and operational dashboards.")
add_rule(doc)

doc.add_heading("1. Product at a glance", level=1)
doc.add_paragraph("ProcureFlow helps farmers choose a procurement centre, confirm eligibility, reserve a token, track queue wait time and follow produce through procurement and payment. Centre staff operate queues; administrators manage capacity, alerts, records and support conversations.")
t = doc.add_table(rows=1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.style = "Light Shading Accent 1"
for c, txt in zip(t.rows[0].cells, ["Farmer", "Centre staff", "Administrator"]):
    shade(c, BLUE); set_cell_text(c, txt, True, "FFFFFF")
row = t.add_row().cells
for c, txt in zip(row, ["Book and track visits; contact support", "Advance queues and procurement stages", "Manage schedules, records, alerts and chats"]): set_cell_text(c, txt)
doc.add_paragraph("Demo accounts", style="Heading 2")
for item in ["Farmer — farmer@demo.in / farmer123", "Centre staff — staff@demo.in / staff123", "Administrator — admin@demo.in / admin123"]: bullet(doc, item)
doc.add_paragraph("No ChatGPT account is required. ProcureFlow uses its own demonstration role sign-in.")
add_picture(doc, "02_access.png", "Role-based ProcureFlow access screen")

doc.add_heading("2. Farmer journey", level=1)
doc.add_heading("2.1 Choose the right centre", level=2)
doc.add_paragraph("Open the Overview page and compare nearby centres by travel distance, available slots, queue size and estimated wait. The recommendation favours the least-congested practical option.")
add_picture(doc, "04_centres.png", "Centre recommendation and congestion comparison")
doc.add_heading("2.2 Book a procurement token", level=2)
steps = ["Open Schedules and select Check eligibility.", "Choose the procurement centre and preferred time window.", "Enter the quantity in quintals.", "Confirm government identity, verified bank account and land/tenancy documentation.", "Press Confirm and book token. Your token, position and ETA appear on the dashboard."]
for s in steps: doc.add_paragraph(s, style="List Number")
add_picture(doc, "05_booking.png", "Eligibility and token-booking flow")
doc.add_heading("2.3 Queue and procurement journey", level=2)
doc.add_paragraph("The live journey is: Booked → Arrived → Gate check-in → Weighing → Quality check → Accepted → Payment processing → Payment completed. Wait time changes as centre staff progress the queue.")
add_picture(doc, "06_journey.png", "Farmer-facing procurement status and estimated wait")
doc.add_heading("2.4 Missed-slot recovery", level=2)
doc.add_paragraph("From the active token, select I missed my slot. Choose the next suitable centre/time and confirm rebooking. The previous visit is replaced with the selected recovery option.")

doc.add_heading("3. ProcureBot automated assistant", level=1)
doc.add_paragraph("Open Support and keep ProcureBot help selected. The rule-based assistant answers common questions instantly without an external AI account. It covers token booking, eligibility documents, queue ETA, missed slots, payment status, regional languages and escalation to an admin.")
bullet(doc, "Use a quick-question chip for the fastest demonstration.")
bullet(doc, "Type a natural short question such as “What documents are needed?”")
bullet(doc, "For an individual case, switch to Live admin chat.")

doc.add_heading("4. Live admin chat", level=1)
doc.add_paragraph("Farmers and centre staff can contact the administrator directly. The conversation displays the participant’s name, stores the history in the hosted database and places media in managed object storage.")
doc.add_heading("4.1 Send text and media", level=2)
for s in ["Select Live admin chat.", "Type a message, or press ＋ to attach a photo, short video or audio file.", "For a voice note, press ●, allow microphone access, speak, then press ● again.", "Press Send. Attachments are limited to 12 MB for reliable mobile use."]: doc.add_paragraph(s, style="List Number")
doc.add_heading("4.2 Read receipts, editing and deletion", level=2)
bullet(doc, "Sent ✓ means the message reached the service.")
bullet(doc, "Seen ✓✓ means the other role opened the conversation.")
bullet(doc, "Edit changes your own text and marks it as edited.")
bullet(doc, "Delete removes message content and its attachment; an admin may moderate any message.")
bullet(doc, "A short notification chime plays for new incoming messages after browser interaction.")
doc.add_heading("4.3 Administrator workflow", level=2)
doc.add_paragraph("Sign in as Administrator, open Live chat, select a named conversation in the left panel and reply. Unread badges identify pending requests. Opening a conversation issues the read receipt.")

doc.add_heading("5. Centre staff and administrator operations", level=1)
doc.add_heading("5.1 Centre staff", level=2)
doc.add_paragraph("The centre dashboard shows green/yellow/red load indicators, expected arrivals and the active queue. Use the action button to progress the selected token through every procurement stage.")
add_picture(doc, "08_staff.png", "Centre load and live queue workflow")
doc.add_heading("5.2 Administrator", level=2)
doc.add_paragraph("The administrator monitors capacity and congestion, creates schedules, explores 100,000 deterministic demonstration records, sends farmer alerts and manages live support.")
add_picture(doc, "09_admin.png", "Administrative overview")
add_picture(doc, "10_records.png", "Filterable and sortable procurement records")

doc.add_heading("6. Installable Android app", level=1)
doc.add_paragraph("ProcureFlow is a Progressive Web App (PWA), so the same tested website can be installed without maintaining a separate APK during the hackathon.")
for s in ["Open the public HTTPS URL in Chrome on Android.", "Select Install app in ProcureFlow when available, or Chrome menu → Add to Home screen.", "Confirm installation and launch ProcureFlow from its home-screen icon."]: doc.add_paragraph(s, style="List Number")
doc.add_paragraph("Core screens are cached for fast reopening. Live queue data, chats, media and alerts require connectivity.")

doc.add_heading("7. Languages and accessibility", level=1)
doc.add_paragraph("Use the language selector in the top bar to switch between English, Hindi and Marathi. The responsive layout uses large touch targets, readable status colours plus labels, and mobile bottom navigation.")
add_picture(doc, "11_language.png", "Regional-language interface")

doc.add_heading("8. Alerts and notifications", level=1)
doc.add_paragraph("Messages contains booking, queue, delay, procurement and payment updates. Enable phone alerts grants browser-notification permission. A custom alert appears in-app and on the current device. The preset Twilio button demonstrates trial SMS delivery; production custom SMS should use approved sender/template registration and applicable compliance controls.")
add_picture(doc, "07_alerts.png", "Notification centre and custom farmer alert")

doc.add_heading("9. Troubleshooting", level=1)
troubles = [
    ("Chat does not load", "Check connectivity, refresh, and sign in again."),
    ("Microphone is blocked", "Open browser site permissions and allow microphone access."),
    ("Install button is missing", "Use Chrome menu → Add to Home screen."),
    ("No sound", "Interact with the page once and check the device mute setting."),
    ("Media is rejected", "Use a photo/video/audio file below 12 MB."),
    ("Language remains English", "Select Hindi or Marathi again, then refresh the current route."),
]
t = doc.add_table(rows=1, cols=2); t.style = "Light Shading Accent 1"
for c, txt in zip(t.rows[0].cells, ["Issue", "Resolution"]): shade(c, BLUE); set_cell_text(c, txt, True, "FFFFFF")
for issue, fix in troubles:
    cells = t.add_row().cells; set_cell_text(cells[0], issue, True); set_cell_text(cells[1], fix)

doc.add_heading("10. Demo and production boundaries", level=1)
doc.add_paragraph("This is a hackathon MVP with deterministic sample procurement data and demonstration identities. The hosted chat does persist conversations and media, but a real rollout must integrate authoritative farmer authentication, role authorization, consent and retention policy, attachment malware scanning, audit logs, rate limiting, backups/restore testing, accessibility review and approved messaging templates.")
doc.add_paragraph("Recommended 4-minute demonstration", style="Heading 2")
for s in ["Farmer: compare centres and book a token.", "Staff: progress the queue and show the ETA/journey update.", "Farmer: ask ProcureBot, then send a photo or voice note to admin.", "Admin: open the named conversation, reply and demonstrate Seen ✓✓.", "Install the PWA and finish with multilingual and records views."]: doc.add_paragraph(s, style="List Number")

doc.save(OUT / "ProcureFlow_User_Manual.docx")
print(OUT / "ProcureFlow_User_Manual.docx")
